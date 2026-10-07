"""Port of src/lib/users-import.ts (Excel parse + row validation)."""

import datetime
import re
from dataclasses import asdict, dataclass
from typing import Optional

from openpyxl import load_workbook

IMPORT_MAX_ROWS = 500
IMPORT_MAX_BYTES = 2 * 1024 * 1024

ROLE_ALIAS = {
    "admin": "admin",
    "dosen": "lecturer",
    "lecturer": "lecturer",
    "mahasiswa": "student",
    "student": "student",
}

STATUS_ALIAS = {
    "aktif": "active",
    "active": "active",
    "nonaktif": "inactive",
    "tidak aktif": "inactive",
    "inactive": "inactive",
}

EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


@dataclass
class ImportRow:
    line: int
    name: str
    email: str
    role: str
    status: str
    error: Optional[str]
    skip: Optional[str]

    def as_dict(self):
        return asdict(self)


def normalize_row(line, raw):
    name = re.sub(r"\s+", " ", (raw.get("name") or "").strip())
    email = (raw.get("email") or "").strip().lower()
    role_raw = (raw.get("role") or "").strip().lower()
    status_raw = (raw.get("status") or "").strip().lower()

    error = None
    if not name:
        error = "Nama wajib diisi."
    elif len(name) > 120:
        error = "Nama maksimal 120 karakter."
    elif not EMAIL_RE.match(email):
        error = "Email tidak valid."

    role = ROLE_ALIAS.get(role_raw, "")
    if not error and not role:
        if role_raw == "":
            role = "student"
        else:
            error = (
                f'Peran tidak valid: "{raw.get("role") or ""}". '
                "Gunakan Admin, Dosen, atau Mahasiswa."
            )

    status = STATUS_ALIAS.get(status_raw, "")
    if not error and not status:
        if status_raw == "":
            status = "active"
        else:
            error = (
                f'Status tidak valid: "{raw.get("status") or ""}". '
                "Gunakan Aktif atau Nonaktif."
            )

    return ImportRow(line, name, email, role, status, error, None)


def detect_in_file_dups(rows):
    seen = set()
    for row in rows:
        if row.error:
            continue
        if row.email in seen:
            row.skip = "Email kembar di file ini."
        else:
            seen.add(row.email)


def mark_existing_emails(rows):
    emails = list(
        {r.email for r in rows if not r.error and not r.skip}
    )
    if not emails:
        return
    from accounts.models import User

    existing = {
        e.lower()
        for e in User.objects.filter(email__in=emails).values_list(
            "email", flat=True
        )
    }
    for row in rows:
        if not row.error and not row.skip and row.email in existing:
            row.skip = "Email sudah terdaftar."


def summarize(rows):
    importable = skipped = errors = 0
    for row in rows:
        if row.error:
            errors += 1
        elif row.skip:
            skipped += 1
        else:
            importable += 1
    return {
        "total": len(rows),
        "importable": importable,
        "skipped": skipped,
        "errors": errors,
    }


def sanitize_client_rows(data):
    if not isinstance(data, list) or len(data) == 0:
        return None
    if len(data) > IMPORT_MAX_ROWS:
        return None
    rows = []
    for i, item in enumerate(data):
        if not isinstance(item, dict):
            return None
        for key in ("name", "email", "role", "status"):
            if not isinstance(item.get(key), str):
                return None
        line = item.get("line")
        if isinstance(line, bool) or not isinstance(line, (int, float)):
            line = i + 1
        rows.append(
            normalize_row(
                line,
                {
                    "name": item["name"],
                    "email": item["email"],
                    "role": item["role"],
                    "status": item["status"],
                },
            )
        )
    detect_in_file_dups(rows)
    return rows


def _js_string(value):
    """JS String() semantics for cell values."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value:  # NaN
            return "NaN"
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime.datetime):
        if value.microsecond:
            base = value.strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3]
        else:
            base = value.strftime("%Y-%m-%dT%H:%M:%S.000")
        return base + "Z" if not value.tzinfo else value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, datetime.time):
        return value.isoformat()
    return str(value)


def cell_text(value):
    if value is None:
        return ""
    return _js_string(value)


def header_key(text):
    v = (text or "").strip().lower()
    if v in ("nama", "nama lengkap", "name"):
        return "name"
    if v in ("email", "e-mail", "surel"):
        return "email"
    if v in ("peran", "role"):
        return "role"
    if v == "status":
        return "status"
    return ""


def parse_users_workbook(data):
    try:
        import io

        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception:
        return [], ["Berkas tidak valid. Pastikan berkas berformat .xlsx."]
    if not workbook.worksheets:
        return [], ["Berkas tidak berisi data. Gunakan template yang disediakan."]
    sheet = workbook.worksheets[0]
    max_row = sheet.max_row or 1
    if max_row < 2:
        return [], ["Berkas tidak berisi data. Gunakan template yang disediakan."]

    header_at = -1
    cols = None
    for r in range(1, min(max_row, 20) + 1):
        found = {"name": -1, "email": -1, "role": -1, "status": -1}
        for cell in sheet[r]:
            if cell.value is None:
                continue
            key = header_key(cell_text(cell.value))
            if key:
                found[key] = cell.column
        if found["name"] != -1 and found["email"] != -1:
            header_at = r
            cols = found
            break
    if cols is None or header_at == -1:
        return [], [
            "Baris judul tidak ditemukan. Kolom Nama dan Email wajib ada. "
            "Gunakan template yang disediakan."
        ]

    rows = []
    file_errors = []
    for r in range(header_at + 1, max_row + 1):
        if len(rows) >= IMPORT_MAX_ROWS:
            file_errors.append(
                f"Maksimal {IMPORT_MAX_ROWS} baris per impor. "
                "Baris berikutnya tidak dibaca."
            )
            break
        row = sheet[r]

        def cell_at(col):
            if col < 1 or col > len(row):
                return ""
            return cell_text(row[col - 1].value)

        raw = {
            "name": cell_at(cols["name"]),
            "email": cell_at(cols["email"]),
            "role": cell_at(cols["role"]) if cols["role"] > 0 else "",
            "status": cell_at(cols["status"]) if cols["status"] > 0 else "",
        }
        if not raw["name"].strip() and not raw["email"].strip():
            continue
        rows.append(normalize_row(r, raw))
    if not rows:
        file_errors.append("Tidak ada baris data.")
    detect_in_file_dups(rows)
    return rows, file_errors
